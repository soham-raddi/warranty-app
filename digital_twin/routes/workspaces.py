import io
import os
import re
from datetime import datetime
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from flask import Blueprint, g, jsonify, redirect, render_template, request, Response, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db import (
    accept_invite_token,
    create_invite,
    create_workspace,
    delete_invite,
    delete_user_account,
    delete_workspace,
    export_workspace_rows,
    get_all_settings,
    get_invite_by_token,
    get_membership_role,
    get_user_workspaces,
    get_workspace,
    list_invites,
    list_workspace_members,
    owned_workspaces_with_other_members,
    parse_amount,
    regenerate_invite_token,
    remove_workspace_member,
    rename_workspace,
    set_setting_value,
    transfer_ownership,
)
from ..services import emails, storage
from ..services.ratelimit import limit
from .auth import is_dev_login_allowed, is_google_configured, role_required

workspaces_bp = Blueprint('workspaces', __name__)

EMAIL_PATTERN = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
MAX_NAME_LENGTH = 60
# Keeps one account from creating workspaces without limit on a public install.
MAX_OWNED_WORKSPACES = 5


def new_workspaces_allowed():
    """ALLOW_NEW_WORKSPACES=false makes the install invite-only (people can only join existing families)."""
    return os.getenv('ALLOW_NEW_WORKSPACES', 'true').lower() in {'1', 'true', 'yes'}


def _clean_name(raw):
    name = re.sub(r'\s+', ' ', str(raw or '')).strip()
    return name[:MAX_NAME_LENGTH]


def valid_timezone(value):
    """The IANA time zone name if the server knows it (e.g. Asia/Kuala_Lumpur), else None."""
    try:
        ZoneInfo(str(value or ''))
        return str(value)
    except (ZoneInfoNotFoundError, ValueError):
        return None


@workspaces_bp.app_template_filter('nice_date')
def nice_date(value):
    """'2026-10-21T18:30:00' -> 'Oct 21, 2026', matching how dates appear elsewhere in the app."""
    try:
        return datetime.fromisoformat(str(value)).strftime('%b %d, %Y')
    except (TypeError, ValueError):
        return value or ''


def _invite_link(token):
    return url_for('workspaces.join_page', token=token, _external=True)


def _email_invite(email, link, expires_at):
    """Emails an invite link. Returns None when sent, or an error message to show instead."""
    if not emails.EMAIL_ENABLED:
        return "Email isn't set up on this site, so share the link yourself."
    inviter = g.current_user.get('name') or g.current_user['email']
    workspace = get_workspace(g.workspace_id) or {'name': 'their workspace'}
    subject, html_body, text_body = emails.invite_email(inviter, workspace['name'], link, nice_date(expires_at))
    try:
        emails.send_email(email, subject, html_body, text_body)
    except emails.EmailError as error:
        return f'{error} You can still share the link below yourself.'
    return None


def _default_workspace_name(user):
    first_name = (user.get('name') or user['email'].split('@')[0]).split()[0]
    return f"{first_name}'s Family"


@workspaces_bp.route('/welcome')
def welcome_page():
    """Shown to signed-in people who do not belong to any workspace yet."""
    user = g.current_user
    if get_user_workspaces(user['id']):
        return redirect(url_for('pages.dashboard_page'))
    return render_template(
        'welcome.html',
        user=user,
        can_create=new_workspaces_allowed(),
        suggested_name=_default_workspace_name(user),
    )


@workspaces_bp.route('/api/workspaces', methods=['GET'])
def list_my_workspaces():
    return jsonify({
        'workspaces': get_user_workspaces(g.current_user['id']),
        'current_workspace_id': session.get('workspace_id'),
        'can_create': new_workspaces_allowed(),
    })


@workspaces_bp.route('/api/workspaces', methods=['POST'])
@limit('create_workspace', daily=10)
def create_my_workspace():
    if not new_workspaces_allowed():
        return jsonify({'error': 'New workspaces are disabled on this site. Ask a family owner to invite you.'}), 403

    user = g.current_user
    owned = [w for w in get_user_workspaces(user['id']) if w['role'] == 'owner']
    if len(owned) >= MAX_OWNED_WORKSPACES:
        return jsonify({'error': f'You can own at most {MAX_OWNED_WORKSPACES} workspaces.'}), 400

    payload = request.json or {}
    name = _clean_name(payload.get('name')) or _default_workspace_name(user)
    workspace_id = create_workspace(name, user['id'])
    timezone = valid_timezone(payload.get('timezone'))
    if timezone:
        set_setting_value(workspace_id, 'timezone', timezone)  # start in the creator's time zone, not UTC
    session['workspace_id'] = workspace_id
    return jsonify({'workspace': {'id': workspace_id, 'name': name, 'role': 'owner'}}), 201


@workspaces_bp.route('/api/workspaces/switch', methods=['POST'])
def switch_workspace():
    payload = request.json or {}
    try:
        workspace_id = int(payload.get('workspace_id'))
    except (TypeError, ValueError):
        return jsonify({'error': 'Choose a workspace.'}), 400

    if not get_membership_role(workspace_id, g.current_user['id']):
        return jsonify({'error': 'You are not a member of that workspace.'}), 403

    session['workspace_id'] = workspace_id
    return jsonify({'message': 'Switched workspace.', 'workspace_id': workspace_id})


@workspaces_bp.route('/api/workspace', methods=['PATCH'])
@role_required('owner')
def rename_current_workspace():
    name = _clean_name((request.json or {}).get('name'))
    if not name:
        return jsonify({'error': 'Enter a workspace name.'}), 400
    rename_workspace(g.workspace_id, name)
    return jsonify({'message': 'Workspace renamed.', 'workspace': {'id': g.workspace_id, 'name': name}})


@workspaces_bp.route('/api/workspace/members', methods=['GET'])
def workspace_members():
    can_manage = g.role in ('owner', 'admin')
    return jsonify({
        'workspace': g.workspace,
        'members': list_workspace_members(g.workspace_id),
        'invites': list_invites(g.workspace_id) if can_manage else [],
        'current_user_id': g.current_user['id'],
        'current_role': g.role,
        'can_manage': can_manage,
        'can_add_admins': g.role == 'owner',
    })


@workspaces_bp.route('/api/workspace/invites', methods=['POST'])
@role_required('owner', 'admin')
@limit('invite', per_workspace=True, hourly=30, daily=100)
def invite_member():
    payload = request.json or {}
    email = str(payload.get('email', '')).strip().lower()
    role = str(payload.get('role', 'member')).strip().lower()

    if not EMAIL_PATTERN.match(email):
        return jsonify({'error': 'Enter a valid email address.'}), 400
    if role not in ('member', 'admin'):
        role = 'member'
    if role == 'admin' and g.role != 'owner':
        return jsonify({'error': 'Only the owner can invite admins.'}), 403

    invite, error = create_invite(g.workspace_id, email, role, g.current_user['id'])
    if error:
        return jsonify({'error': error}), 400
    token = invite.pop('token')
    link = _invite_link(token)
    email_error = None
    if payload.get('send_email'):
        email_error = _email_invite(email, link, invite['expires_at'])
    if payload.get('send_email') and not email_error:
        message = f'Invite emailed to {email}. You can also share the link below.'
    else:
        message = email_error or f'Invite created for {email}. Send them the link below.'
    return jsonify({
        'message': message,
        'invite': invite,
        'link': link,
        'email_sent': bool(payload.get('send_email')) and not email_error,
    }), 201


@workspaces_bp.route('/api/workspace/invites/<int:invite_id>/link', methods=['POST'])
@role_required('owner', 'admin')
@limit('invite', per_workspace=True, hourly=30, daily=100)
def new_invite_link(invite_id):
    # Only a hash of each link is stored, so a lost link is replaced rather than shown again.
    result = regenerate_invite_token(g.workspace_id, invite_id)
    if not result:
        return jsonify({'error': 'That invite was not found. It may have been used or revoked.'}), 404
    token, expires_at = result
    return jsonify({'link': _invite_link(token), 'expires_at': expires_at,
                    'message': 'New invite link created. The previous link no longer works.'})


@workspaces_bp.route('/api/workspace/invites/<int:invite_id>/email', methods=['POST'])
@role_required('owner', 'admin')
@limit('invite', per_workspace=True, hourly=30, daily=100)
def email_invite_again(invite_id):
    invite = next((i for i in list_invites(g.workspace_id) if i['id'] == invite_id), None)
    result = regenerate_invite_token(g.workspace_id, invite_id) if invite else None
    if not result:
        return jsonify({'error': 'That invite was not found. It may have been used or revoked.'}), 404
    token, expires_at = result
    link = _invite_link(token)
    email_error = _email_invite(invite['email'], link, expires_at)
    if email_error:
        return jsonify({'error': email_error, 'link': link, 'expires_at': expires_at, 'email': invite['email']}), 502
    return jsonify({'message': f"Invite emailed to {invite['email']} again. Older links no longer work.",
                    'link': link, 'expires_at': expires_at})


@workspaces_bp.route('/api/workspace/invites/<int:invite_id>', methods=['DELETE'])
@role_required('owner', 'admin')
def revoke_invite(invite_id):
    if not delete_invite(g.workspace_id, invite_id):
        return jsonify({'error': 'That invite was not found. It may have been used or revoked.'}), 404
    return jsonify({'message': 'Invite revoked.'})


@workspaces_bp.route('/api/workspace/members/<int:user_id>', methods=['DELETE'])
@role_required('owner', 'admin')
def remove_member(user_id):
    if user_id == g.current_user['id']:
        return jsonify({'error': 'Use "Leave workspace" to remove yourself.'}), 400

    target_role = get_membership_role(g.workspace_id, user_id)
    if target_role is None:
        return jsonify({'error': 'This person is not a member of the workspace.'}), 404
    if g.role == 'admin' and target_role != 'member':
        return jsonify({'error': 'Admins can only remove members.'}), 403

    removed, message = remove_workspace_member(g.workspace_id, user_id)
    if not removed:
        return jsonify({'error': message}), 400
    return jsonify({'message': message})


@workspaces_bp.route('/api/workspace/leave', methods=['POST'])
def leave_workspace():
    if g.role == 'owner':
        return jsonify({'error': 'The owner cannot leave their own workspace. Transfer ownership first.'}), 400

    remove_workspace_member(g.workspace_id, g.current_user['id'])
    remaining = get_user_workspaces(g.current_user['id'])
    session['workspace_id'] = remaining[0]['id'] if remaining else None
    return jsonify({'message': 'You left the workspace.', 'redirect': url_for('pages.dashboard_page') if remaining else url_for('workspaces.welcome_page')})


@workspaces_bp.route('/join/<token>', methods=['GET'])
def join_page(token):
    """Landing page for an invite link. Viewing it never joins anyone; that needs the Join button (POST)."""
    invite = get_invite_by_token(token)
    user = g.get('current_user')
    already_member = bool(invite and user and get_membership_role(invite['workspace_id'], user['id']))
    if invite and invite['accepted'] and not already_member:
        invite = None  # a used link means nothing to anyone else

    return render_template(
        'join.html',
        invite=invite,
        user=user,
        token=token,
        already_member=already_member,
        google_configured=is_google_configured(),
        allow_dev_login=is_dev_login_allowed(),
    ), (200 if invite else 404)


@workspaces_bp.route('/join/<token>', methods=['POST'])
def accept_invite(token):
    user = g.get('current_user')
    if not user:
        return redirect(url_for('auth.login_page', next=url_for('workspaces.join_page', token=token)))

    workspace_id, error = accept_invite_token(token, user)
    if error:
        invite = get_invite_by_token(token)
        return render_template('join.html', invite=invite, user=user, token=token, error=error,
                               already_member=False, google_configured=is_google_configured(),
                               allow_dev_login=is_dev_login_allowed()), 400

    session['workspace_id'] = workspace_id
    return redirect(url_for('pages.dashboard_page'))


@workspaces_bp.route('/api/workspace/transfer', methods=['POST'])
@role_required('owner')
def transfer_workspace():
    try:
        new_owner_id = int((request.json or {}).get('user_id'))
    except (TypeError, ValueError):
        return jsonify({'error': 'Choose a member to make the owner.'}), 400
    error = transfer_ownership(g.workspace_id, g.current_user['id'], new_owner_id)
    if error:
        return jsonify({'error': error}), 400
    return jsonify({'message': 'Ownership transferred. You are now an admin of this workspace.'})


@workspaces_bp.route('/api/workspace', methods=['DELETE'])
@role_required('owner')
def delete_current_workspace():
    """Permanently deletes the workspace, its items, receipts, settings, invites and chats."""
    confirm = str((request.json or {}).get('confirm_name', '')).strip()
    if confirm != (g.workspace or {}).get('name'):
        return jsonify({'error': 'Type the workspace name exactly to confirm.'}), 400
    for ref in delete_workspace(g.workspace_id):
        storage.delete(ref)
    session.pop('workspace_id', None)
    remaining = get_user_workspaces(g.current_user['id'])
    return jsonify({'message': 'Workspace deleted.',
                    'redirect': url_for('pages.dashboard_page') if remaining else url_for('workspaces.welcome_page')})


@workspaces_bp.route('/api/workspace/export.xlsx', methods=['GET'])
def export_workspace():
    """Inventory and service history as an Excel workbook, one sheet each (receipt files stay downloadable)."""
    items, history = export_workspace_rows(g.workspace_id)
    settings = get_all_settings(g.workspace_id)
    home = settings.get('currency_code') or 'INR'
    default_days = settings.get('warranty_period_days') or '365'

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Inventory'
    inventory_rows = []
    for item in items:
        months = item.get('warranty_months')
        inventory_rows.append([
            item.get('item_name'), item.get('brand'), item.get('model_number'), item.get('serial_number'),
            item.get('category'), item.get('store_name'), item.get('invoice_number'),
            _as_date(item.get('date_of_purchase')), parse_amount(item.get('total_amount')), item.get('currency'),
            item.get('home_amount'),
            f'{months} months' if months else f'{default_days} days (default)',
            _as_date(item.get('warranty_expires_on')), item.get('warranty_status'),
            'Yes' if int(item.get('has_warranty_card') or 0) else 'No',
            item.get('warranty_info'), item.get('return_policy'),
        ])
    _write_sheet(sheet, [
        ('Item', 40), ('Brand', 16), ('Model', 18), ('Serial number', 18), ('Category', 20), ('Store', 26),
        ('Invoice number', 18), ('Purchase date', 14, 'date'), ('Amount', 13, 'money'), ('Currency', 10),
        (f'Amount in {home}', 16, 'money'), ('Warranty', 18), ('Warranty ends', 14, 'date'), ('Status', 14),
        ('Warranty card', 14), ('Warranty notes', 40, 'wrap'), ('Return policy', 40, 'wrap'),
    ], inventory_rows, total_column=11 if inventory_rows else None,
        empty_message='No items yet.')

    history_rows = [[
        entry.get('item_name'), _as_date(entry.get('service_date')), entry.get('service_type'), entry.get('vendor'),
        parse_amount(entry.get('cost')), entry.get('currency'), entry.get('notes'), _as_date(entry.get('next_service_date')),
    ] for entry in history]
    _write_sheet(workbook.create_sheet('Service history'), [
        ('Item', 40), ('Date', 14, 'date'), ('Service', 26), ('Vendor', 22), ('Cost', 13, 'money'), ('Currency', 10),
        ('Notes', 40, 'wrap'), ('Next service', 14, 'date'),
    ], history_rows, empty_message='No service records yet.')

    buffer = io.BytesIO()
    workbook.save(buffer)
    safe_name = re.sub(r'[^A-Za-z0-9_-]+', '-', (g.workspace or {}).get('name', 'workspace')).strip('-') or 'workspace'
    filename = f"{safe_name}-inventory-{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return Response(buffer.getvalue(), headers={'Content-Disposition': f'attachment; filename="{filename}"'},
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


HEADER_FILL = PatternFill('solid', fgColor='1D4ED8')
MONEY_FORMAT = '#,##0.00'
DATE_FORMAT = 'dd mmm yyyy'


def _write_sheet(sheet, columns, rows, total_column=None, empty_message=''):
    """Writes a table: bold header that stays visible and can be filtered, sized columns, real dates and numbers."""
    for index, (title, width, *kind) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center')
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    if not rows:
        sheet.cell(row=2, column=1, value=empty_message).font = Font(italic=True, color='64748B')
        return
    for row_number, row in enumerate(rows, start=2):
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=index)
            kind = columns[index - 1][2] if len(columns[index - 1]) > 2 else None
            if isinstance(value, str):
                # Text from receipts is always text, never a formula, even if it starts with "=".
                cell.value = value
                cell.data_type = 's'
            else:
                cell.value = value
            if kind == 'money':
                cell.number_format = MONEY_FORMAT
            elif kind == 'date':
                cell.number_format = DATE_FORMAT
            elif kind == 'wrap':
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    last = len(rows) + 1
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{last}'
    if total_column:
        letter = get_column_letter(total_column)
        label = sheet.cell(row=last + 1, column=total_column - 1, value='Total')
        total = sheet.cell(row=last + 1, column=total_column, value=f'=SUM({letter}2:{letter}{last})')
        label.font = total.font = Font(bold=True)
        total.number_format = MONEY_FORMAT


def _as_date(value):
    try:
        return datetime.strptime(str(value), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None



@workspaces_bp.route('/api/account', methods=['DELETE'])
def delete_my_account():
    """Deletes the signed-in person's account, workspaces they own alone, and their chats."""
    user = g.current_user
    if str((request.json or {}).get('confirm_email', '')).strip().lower() != user['email'].lower():
        return jsonify({'error': 'Type your email address exactly to confirm.'}), 400
    shared = owned_workspaces_with_other_members(user['id'])
    if shared:
        names = ', '.join(w['name'] for w in shared)
        return jsonify({'error': f'You own workspaces other people still use ({names}). '
                                 'Transfer ownership or delete them first.'}), 400
    for ref in delete_user_account(user['id']):
        storage.delete(ref)
    session.clear()
    return jsonify({'message': 'Your account has been deleted.', 'redirect': url_for('auth.login_page')})
